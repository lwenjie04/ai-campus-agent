from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT_PATH = r"d:\VUE3-GDES-Agent\ai-campus-agent\digital-human-case-fixed.docx"


CONTENT = [
    ("数字人智能问答系统项目实践案例", "title"),
    ("一、案例简介", "h1"),
    (
        "本案例选取“数智校答——基于 DeepSeek 与 RAG 的校园数字人问答平台”作为教学案例。该项目来源于学生真实开发实践，面向高校校园服务场景，围绕校园信息分散、查询效率低、服务入口不统一、人工答疑重复度高等现实问题，设计并实现了一套集校园智能问答、知识库检索增强、数字人展示、语音播报、学生社区、管理员审核与知识沉淀于一体的综合型智能服务平台。",
        "body",
    ),
    (
        "系统以前后端分离架构为基础，前端采用 Vue 3、TypeScript、Pinia、Element Plus 等技术实现页面展示与交互，后端采用 Node.js 和 MySQL 提供业务支撑，并结合 DeepSeek 大模型、RAG 检索增强、邮件验证码认证和腾讯云 TTS 语音合成等能力，实现了较完整的智能问答与校园服务流程。该案例兼具技术性、实践性和工程完整性，适合作为软件工程、人工智能应用开发、智能系统设计等课程中的教学案例。通过本案例，学生能够理解真实项目中需求分析、系统设计、模型应用、业务实现与部署落地之间的关系，并进一步认识大模型技术在校园垂直场景中的实际应用价值。",
        "body",
    ),
    ("二、案例内容", "h1"),
    ("2.1 项目背景", "h2"),
    (
        "在高校教学与管理过程中，学生经常需要查询选课安排、补退选规则、考试通知、奖助政策、科研活动、学生事务以及生活服务等信息。然而，这些信息通常分散在学校官网、教务系统、各职能部门页面和附件文件中，学生往往需要在多个平台之间来回查找，不仅耗时，还容易遗漏重要内容。尤其在奖学金申请、转专业办理、课程调整和宿舍服务等高频场景中，学生常常需要依靠个人经验或同学口口相传来补足信息理解上的不足，信息获取成本较高。",
        "body",
    ),
    (
        "随着大语言模型的发展，智能问答系统为高校信息服务提供了新的技术路径。相比传统关键词搜索，大模型具有更强的自然语言理解和表达能力，能够更好地适应学生口语化提问习惯。但如果完全依赖通用大模型生成答案，又容易出现与学校真实制度不一致、回答缺乏依据甚至产生“幻觉”的问题。因此，如何让系统既具备自然语言交互能力，又能保证答案来源于学校真实资料，是本项目的核心问题。",
        "body",
    ),
    (
        "基于这一背景，项目团队提出建设“校园数字人问答平台”，以学校公开通知和相关资料为知识基础，结合大模型问答、检索增强生成、数字人展示和语音播报能力，构建一个既实用又具有展示性的校园智能服务平台。",
        "body",
    ),
    ("2.2 系统主要内容", "h2"),
    (
        "本项目围绕“校园信息查询与智能服务”展开，整体由前端用户交互层、后端业务服务层、知识支撑层、数字人交互层和后台治理层组成。普通用户可以在首页通过自然语言提出问题，系统结合知识库内容返回回答，并展示来源依据；管理员则可以在后台审核帖子、回复和社区知识，保证内容质量并推动知识沉淀。",
        "body",
    ),
    (
        "系统的核心业务链路为：用户输入问题后，后端首先对问题进行处理，并从校园知识库和社区知识中检索出最相关内容，将检索结果作为上下文输入大模型，从而生成更贴近校园场景的回答。前端以聊天式方式呈现结果，同时通过数字人状态切换和语音播报增强交互体验。为进一步提升系统的可持续更新能力，平台还设计了学生社区模块，使学生在交流过程中形成的高质量经验内容，经管理员审核后可以沉淀为社区知识，并反哺问答系统。",
        "body",
    ),
    ("图2-1 系统总体架构图", "caption"),
    ("（此处插入系统总体架构图）", "placeholder"),
    ("2.3 核心功能设计", "h2"),
    ("（1）校园信息智能问答功能", "h2"),
    (
        "该功能是系统的核心服务入口。用户可围绕校园高频场景直接提出问题，如“国家励志奖学金申请条件是什么”“转专业后补退选怎么处理”“宿舍报修如何提交”等，系统根据知识检索结果和大模型能力生成回答，并在页面中展示信息来源。通过这种方式，系统在一定程度上解决了传统检索路径复杂、信息理解成本高的问题。",
        "body",
    ),
    ("（2）RAG 检索增强功能", "h2"),
    ("系统并非简单依赖大模型记忆，而是通过知识库检索增强问答质量，其基本过程可表示为：", "body"),
    ("A = f(Q, R(K, Q))", "formula"),
    (
        "其中，Q 表示用户问题，K 表示知识库，R(K,Q) 表示从知识库中召回的相关内容，A 表示最终生成的回答。该公式说明系统回答建立在知识检索结果基础上，而不是完全依靠模型自由生成。",
        "body",
    ),
    ("为提高召回质量，系统进一步综合考虑关键词匹配、语义相似度和来源权重，其排序思想可表示为：", "body"),
    ("Score = αS_keyword + βS_semantic + γS_source", "formula"),
    (
        "其中，S_keyword 表示关键词匹配得分，S_semantic 表示语义相似得分，S_source 表示知识来源权重；α、β、γ 为对应权重。通过这种方式，系统能够优先选择更相关、更新、更可信的知识内容。",
        "body",
    ),
    ("（3）数字人展示与语音播报功能", "h2"),
    (
        "为了增强系统的展示效果和交互体验，项目在首页设计了数字人展示模块。数字人包含欢迎、待机和讲解三种状态，通过视频切换配合问答流程展示不同场景。当用户进入系统时，数字人以欢迎状态出现；当系统处于等待输入阶段时，数字人保持待机；当系统返回回答并进行语音播报时，数字人切换为讲解状态。通过视频和语音结合，系统从传统文本问答扩展为多模态交互。",
        "body",
    ),
    ("图2-2 首页数字人智能问答界面", "caption"),
    ("（此处插入首页截图）", "placeholder"),
    ("（4）学生社区与知识沉淀功能", "h2"),
    (
        "学生社区模块为用户提供了发帖、回复、浏览、搜索和分类查看等功能，主要服务于奖助、考试、课程安排、校园生活等场景下的经验交流。为了避免社区内容直接影响问答质量，系统引入管理员审核机制，对通过审核的优质内容进一步整理为社区知识条目，作为官方知识的补充来源。这样，系统形成了“问答服务—社区补充—后台审核—知识沉淀—问答增强”的闭环结构。",
        "body",
    ),
    ("图2-3 学生社区界面", "caption"),
    ("（此处插入学生社区截图）", "placeholder"),
    ("（5）管理员审核与治理功能", "h2"),
    (
        "管理员审核中心主要负责对帖子、回复和社区知识条目进行审核与管理。系统在后台展示待审核帖子、待审核回复、待入库候选和待审核知识等内容，管理员可进行通过、驳回和知识生成操作。该模块使平台具备基本的内容治理能力，保证知识沉淀过程更加规范。",
        "body",
    ),
    ("图2-4 管理员审核中心界面", "caption"),
    ("（此处插入管理员审核页面截图）", "placeholder"),
    ("2.4 项目实施过程与主要问题", "h2"),
    (
        "在项目实施过程中，团队首先完成了需求分析和模块划分，随后逐步实现首页问答、知识库支撑、数字人展示、学生社区、后台审核和登录注册等功能。在知识建设阶段，团队对学校公开通知和相关资料进行了筛选、整理和分类，作为问答系统的知识支撑基础。在系统开发阶段，前端重点完成页面交互与数字人展示，后端重点完成问答接口、用户认证、邮件发送、TTS 语音合成和社区审核逻辑。",
        "body",
    ),
    (
        "项目实施中主要遇到了三类问题。第一，知识内容分散且格式不统一，导致知识整理和结构化处理成本较高；第二，单纯调用通用大模型时，回答与学校实际信息不完全一致，因此必须引入知识检索增强机制；第三，数字人视频切换、语音播报、前后端联调和部署运行等工程问题较多，需要通过多轮调试与优化逐步解决。",
        "body",
    ),
    ("2.5 项目成效", "h2"),
    (
        "经过持续开发和优化，项目已完成校园数字人智能问答系统的主要功能建设，实现了智能问答、知识来源展示、数字人展示、学生社区、管理员审核、邮箱验证码注册和网站部署等内容，具备实际访问和展示条件。项目从概念设计逐步发展为可运行、可演示、可持续优化的平台系统，既体现了大模型技术在校园场景中的应用价值，也反映了大学生在项目式开发中的工程实践能力。",
        "body",
    ),
    ("三、教学目标", "h1"),
    ("3.1 知识目标", "h2"),
    (
        "通过本案例的学习，使学生掌握校园数字人智能问答系统的基本构成与实现思路，理解前后端分离架构、RAG 检索增强、大模型调用、数字人展示、邮件验证码认证和管理员审核等技术在真实项目中的应用方式，理解知识支撑、业务流程与系统治理之间的关系。",
        "body",
    ),
    ("3.2 能力目标", "h2"),
    (
        "培养学生从真实场景中提炼需求并完成功能设计的能力，提高学生在前端开发、后端接口设计、数据库存储、智能服务集成、系统联调与部署运行等方面的综合工程实践能力，增强学生分析真实项目案例和归纳系统实现逻辑的能力。",
        "body",
    ),
    ("3.3 素质目标", "h2"),
    (
        "引导学生树立面向真实问题开展项目开发的意识，增强工程思维、系统思维和团队协作意识，提升学生在项目表达、技术论证和成果展示方面的综合素养。",
        "body",
    ),
    ("四、教学实施", "h1"),
    ("4.1 教学实施要求", "h2"),
    (
        "本案例适合用于软件工程、人工智能应用开发、创新创业实践和智能系统设计等课程教学。教学中应坚持“问题导向、技术拆解、系统分析与案例反思”相结合的原则，引导学生从应用场景出发，理解系统需求、功能结构、关键技术和工程实现之间的关系。教师可结合架构图、页面截图和检索评分公式开展讲解，增强案例的直观性。",
        "body",
    ),
    ("4.2 教学实施步骤", "h2"),
    (
        "第一阶段为案例导入。教师可从校园信息分散、检索低效、人工答疑重复等问题出发，引出数字人智能问答平台的建设背景。第二阶段为需求分析。组织学生分析系统需要解决的核心问题，拆解问答、社区、审核、数字人展示和认证等模块。第三阶段为技术分析。讲解系统采用的关键技术，包括 DeepSeek 大模型、RAG 检索增强、数字人状态切换、TTS 语音播报、数据库支撑和邮件验证等。第四阶段为界面与流程分析。结合首页、社区页和管理员页截图分析页面设计与业务流程之间的关系。第五阶段为讨论与反思。围绕问答准确性、知识更新机制、数字人交互自然度和系统可扩展性等问题组织学生讨论。第六阶段为总结提升。教师对案例中的技术路线、系统设计和应用价值进行归纳，并引导学生思考其未来优化方向。",
        "body",
    ),
    ("4.3 教学方法", "h2"),
    (
        "本案例建议采用案例教学法、任务驱动法和讨论式教学法相结合的方式实施。教师可按照“案例背景—需求分析—系统拆解—技术分析—课堂讨论—总结提升”的流程组织教学，帮助学生把抽象技术知识与具体应用场景联系起来。",
        "body",
    ),
    ("五、总结", "h1"),
    (
        "本案例围绕“数智校答——基于 DeepSeek 与 RAG 的校园数字人问答平台”展开，展示了一个面向高校真实场景的智能服务系统从需求分析、系统设计、功能实现到部署上线的完整过程。项目将校园信息智能问答、知识库检索增强、数字人展示、学生社区、管理员审核和知识沉淀等功能有机结合，体现了大模型技术在校园垂直场景中的综合应用价值。",
        "body",
    ),
    (
        "通过本案例教学，学生不仅能够理解大语言模型与 RAG 技术在实际系统中的作用，还能够从中学习真实项目中的模块划分、工程实现、部署运行与内容治理逻辑。该案例兼具技术性、实践性和教学性，适合作为软件工程和人工智能相关课程中的案例材料。同时，它也提出了进一步思考：未来如何持续优化问答准确性、改进知识更新机制、增强数字人交互自然度，并推动系统在更广泛高校场景中的应用推广。",
        "body",
    ),
]


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3.0)
section.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.font.size = Pt(12)


def set_run_font(run, east="宋体", west="Times New Roman", size=12, bold=False):
    run.font.name = west
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(doc, text, kind="body"):
    p = doc.add_paragraph()
    if kind == "title":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, east="黑体", west="Times New Roman", size=16, bold=True)
    elif kind == "h1":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        set_run_font(run, east="黑体", west="Times New Roman", size=14, bold=True)
    elif kind == "h2":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        set_run_font(run, east="黑体", west="Times New Roman", size=12, bold=True)
    elif kind == "caption":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, east="宋体", west="Times New Roman", size=10.5, bold=False)
    elif kind == "placeholder":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, east="宋体", west="Times New Roman", size=10.5, bold=False)
    elif kind == "formula":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, east="Times New Roman", west="Times New Roman", size=12, bold=False)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        set_run_font(run, east="宋体", west="Times New Roman", size=12, bold=False)
    return p


for text, kind in CONTENT:
    add_paragraph(doc, text, kind)

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
set_run_font(run, east="宋体", west="Times New Roman", size=10.5, bold=False)
fldChar1 = OxmlElement("w:fldChar")
fldChar1.set(qn("w:fldCharType"), "begin")
instrText = OxmlElement("w:instrText")
instrText.set(qn("xml:space"), "preserve")
instrText.text = " PAGE "
fldChar2 = OxmlElement("w:fldChar")
fldChar2.set(qn("w:fldCharType"), "end")
run._r.append(fldChar1)
run._r.append(instrText)
run._r.append(fldChar2)

doc.save(OUT_PATH)
print(OUT_PATH)
